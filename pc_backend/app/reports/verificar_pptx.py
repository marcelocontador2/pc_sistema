from pptx import Presentation
import os
import glob
pasta = r"C:\Users\Marcelo Medeiros\pc_sistema\pc_storage\output_documentos"
arquivos = glob.glob(os.path.join(pasta, "*.pptx"))
if not arquivos:
    print("Nenhum arquivo PPTX encontrado.")
else:
    caminho = max(arquivos, key=os.path.getctime)
    print(f"Verificando: {os.path.basename(caminho)}")
    print("="*70)
    prs = Presentation(caminho)
    slides_para_verificar = [7, 8, 9, 11, 12, 21, 22, 30, 31, 34, 35, 38]
    for i, slide in enumerate(prs.slides):
        if (i + 1) in slides_para_verificar:
            print(f"\n📄 Slide {i+1}:")
            print("-"*50)
            for j, shape in enumerate(slide.shapes):
                if shape.has_text_frame:
                    texto = shape.text_frame.text[:200].replace("\n", " | ")
                    if texto.strip():
                        print(f"  [{j}] '{texto}'")
