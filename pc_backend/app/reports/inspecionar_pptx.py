from pptx import Presentation
import os

caminho_pptx = r"C:\Users\Marcelo Medeiros\pc_sistema\pc_storage\templates\Cofen_Analise_Demonstracoes_Contabeis_1S2026 (10082026).pptx"

if not os.path.exists(caminho_pptx):
    print(f"Arquivo não encontrado: {caminho_pptx}")
else:
    prs = Presentation(caminho_pptx)
    print(f"Total de slides: {len(prs.slides)}")
    print("="*70)

    for i, slide in enumerate(prs.slides):
        print(f"\n📄 Slide {i+1}:")
        print("-"*50)
        for j, shape in enumerate(slide.shapes):
            tipo = shape.shape_type
            nome = shape.name
            texto = ""
            if shape.has_text_frame:
                texto = shape.text_frame.text[:120].replace("\n", " | ")
            print(f"  [{j}] Tipo: {tipo} | Nome: '{nome}' | Texto: '{texto}'")